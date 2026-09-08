"""Checks data/germany.json against the BKA workbooks it was built from.

    python tools/3_bka_fetch.py && python tools/4_germany.py && python tools/verify_germany.py

Same standard as verify_data.py: every published value is re-read out of the
source workbook, so no figure on the Germany panel can drift away from the PKS
without the build failing. The extra thing this suite guards is the promise the
panel makes about itself - that it is never merged with the Eurostat series.
"""
import io
import json
import os

import openpyxl

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(HERE, "data", "raw", "bka")
fails = []


def ck(name, cond, extra=""):
    print(("PASS " if cond else "FAIL ") + name + ("  " + str(extra) if extra else ""))
    if not cond:
        fails.append(name)


with io.open(os.path.join(HERE, "data", "germany.json"), encoding="utf-8") as f:
    d = json.load(f)
with io.open(os.path.join(HERE, "data", "crime.json"), encoding="utf-8") as f:
    eu = json.load(f)

meta = d["meta"]
keys = [c["key"] for c in d["categories"]]

ck("eight offence keys, as documented", len(keys) == 8, keys)
ck("every key has a national series and a Bundesland value",
   all(k in d["national"] for k in keys)
   and all(all(k in v for k in keys) for v in d["laender"]["values"].values()))
ck("sixteen Bundeslaender, the national row held out separately",
   len(d["laender"]["order"]) == 16 and "national_row" in d["laender"],
   len(d["laender"]["order"]))

years = [p["year"] for p in d["national"]["------"]]
ck("the national series runs unbroken from its first year to the reporting year",
   years == list(range(years[0], meta["reporting_year"] + 1)),
   "%d-%d, %d years" % (years[0], years[-1], len(years)))
ck("every category covers the same years",
   all([p["year"] for p in d["national"][k]] == years for k in keys))

# --- nothing may be invented: re-read both workbooks ----------------------
wb = openpyxl.load_workbook(os.path.join(RAW, "ZR-Bund-Faelle.xlsx"), read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
drift = []
seen = 0
for row in ws.iter_rows(values_only=True):
    if row[0] not in keys or not isinstance(row[2], (int, float)):
        continue
    got = next((p for p in d["national"][row[0]] if p["year"] == int(row[2])), None)
    seen += 1
    if (not got or got["cases"] != int(row[3])
            or got["per100k"] != round(float(row[4]), 1)
            or got["clearance"] != round(float(row[9]), 1)):
        drift.append((row[0], row[2], got))
wb.close()
ck("every national value matches the BKA time series workbook", not drift and seen == 8 * len(years),
   str(drift[:2]) + " " + str(seen) + " rows read")

wb = openpyxl.load_workbook(os.path.join(RAW, "LA-F-02-HZ.xlsx"), read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
drift = []
seen = 0
for row in ws.iter_rows(values_only=True):
    if row[0] not in keys or not isinstance(row[2], str) or row[4] is None:
        continue
    name = " ".join(row[2].split())
    got = d["laender"]["values"].get(name) or (d["laender"]["national_row"] if name == "Bundesrepublik Deutschland" else None)
    if got is None:
        continue
    seen += 1
    cell = got[row[0]]
    if (cell["cases"] != int(row[3]) or cell["per100k"] != round(float(row[4]), 1)
            or cell["clearance"] != round(float(row[8]), 1)):
        drift.append((row[0], name, cell))
wb.close()
ck("every Bundesland value matches the BKA Laender workbook", not drift and seen == 8 * 17,
   str(drift[:2]) + " " + str(seen) + " rows read")

# --- the two workbooks have to agree with each other ----------------------
mismatch = [k for k in keys
            if d["laender"]["national_row"][k]["cases"] != d["national"][k][-1]["cases"]]
ck("the Bundesland table's national row equals the time series", not mismatch, mismatch)

sums = [k for k in keys
        if sum(d["laender"]["values"][n][k]["cases"] for n in d["laender"]["order"])
        != d["laender"]["national_row"][k]["cases"]]
ck("the sixteen Bundeslaender add up to the national figure", not sums, sums)

# --- the panel's own promise: separate from Eurostat, never spliced -------
ck("the dataset names what it is not comparable with",
   "crime.json" in meta["not_comparable_with"] and "Eurostat" in meta["not_comparable_with"])
ck("BKA runs ahead of Eurostat, which is why the panel exists",
   meta["reporting_year"] > int(eu["years"][-1]),
   "BKA %d, Eurostat %s" % (meta["reporting_year"], eu["years"][-1]))
ck("no Eurostat ICCS code appears anywhere in the Germany dataset",
   not any(k.startswith("ICCS") for k in keys)
   and "ICCS" not in json.dumps(d["national"]) + json.dumps(d["laender"]))

# --- BKA's reuse condition is source, reporting year AND file version ----
ck("both file versions are recorded, as BKA's reuse terms require",
   all(v.startswith("V") and "erstellt am" in v for v in meta["versions"].values()),
   meta["versions"])
ck("the reuse terms are carried in the dataset",
   "version" in meta["reuse_terms"] and "reporting year" in meta["reuse_terms"])
ck("each source file is recorded with its size and hash",
   len(meta["files"]) == 2 and all(len(x["sha256"]) == 64 and x["bytes"] > 0 for x in meta["files"]))

# --- the additions Eurostat cannot make must actually be present ---------
labels = " ".join(c["label"] for c in d["categories"]).lower()
ck("shoplifting is a category, which is the point of this panel", "shoplifting" in labels)
ck("clearance rates are published, which Eurostat has no field for",
   all(0 <= p["clearance"] <= 100 for k in keys for p in d["national"][k])
   and any(p["clearance"] > 0 for p in d["national"]["*26*00"]))

blob = " ".join(d["limits"]).lower()
ck("the limits say these figures are not comparable with the Eurostat series",
   "not comparable" in blob and "eurostat" in blob)
ck("the limits say a clearance rate is not a recovery rate",
   "not the share of loss recovered" in blob)
ck("the limits say high shoplifting clearance is a recording artefact",
   "recording artefact" in blob)
ck("the limits say the city states read high by construction",
   "commuter" in blob and "denominator" in blob)
ck("the limits say the population basis changed", "census" in blob)
ck("the limits say recorded offences are a floor", "floor" in blob)

print()
if fails:
    print("FAILURES:", len(fails), fails)
    raise SystemExit(1)
print("all Germany checks passed")
